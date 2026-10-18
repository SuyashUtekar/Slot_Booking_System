import os
import re
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from uuid import uuid4

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from sqlalchemy import func, select, update
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, joinedload

from app.config import get_settings
from app.models import EventSettings, ExportOutbox, Registration, RegistrationStatus, SlotStatus, TimeSlot
from app.schemas import RegistrationCreate


def get_event_settings(db: Session) -> EventSettings:
    value = db.get(EventSettings, 1)
    if value is None:
        config = get_settings()
        value = EventSettings(id=1, max_guests_per_registration=config.max_guests_per_registration, timezone=config.event_timezone)
        db.add(value)
        try:
            db.commit()
            db.refresh(value)
        except IntegrityError:
            # Concurrent first requests may both initialize the singleton row.
            # The unique primary key makes one winner; reload that row here.
            db.rollback()
            value = db.get(EventSettings, 1)
            if value is None:
                raise
    return value


def normalize_phone(phone: str) -> str:
    return re.sub(r"[^+0-9]", "", phone)


def public_slot(slot: TimeSlot, now: datetime) -> dict:
    remaining = max(0, slot.capacity - slot.booked_guests)
    starts_at = slot.starts_at if slot.starts_at.tzinfo else slot.starts_at.replace(tzinfo=timezone.utc)
    selectable = slot.status == SlotStatus.OPEN and starts_at > now and remaining > 0
    status = "FULL" if slot.status == SlotStatus.OPEN and remaining == 0 else slot.status
    return {"id": slot.id, "label": slot.label, "starts_at": slot.starts_at, "ends_at": slot.ends_at, "capacity": slot.capacity, "booked_guests": slot.booked_guests, "remaining_guests": remaining, "status": status, "selectable": selectable}


def create_registration(db: Session, payload: RegistrationCreate) -> Registration:
    config = get_event_settings(db)
    if payload.guest_count > config.max_guests_per_registration:
        raise ValueError("GUEST_COUNT_EXCEEDS_LIMIT")
    existing = db.scalar(select(Registration).where(Registration.idempotency_key == payload.idempotency_key).options(joinedload(Registration.slot)))
    if existing:
        return existing
    now = datetime.now(timezone.utc)
    result = db.execute(update(TimeSlot).where(TimeSlot.id == payload.slot_id, TimeSlot.status == SlotStatus.OPEN, TimeSlot.starts_at > now, TimeSlot.booked_guests + payload.guest_count <= TimeSlot.capacity).values(booked_guests=TimeSlot.booked_guests + payload.guest_count, updated_at=func.now()))
    if result.rowcount != 1:
        db.rollback()
        raise LookupError("SLOT_FULL_OR_UNAVAILABLE")
    slot = db.get(TimeSlot, payload.slot_id)
    registration = Registration(confirmation_code=f"VR-{uuid4().hex[:10].upper()}", idempotency_key=payload.idempotency_key, full_name=payload.full_name, phone_number=normalize_phone(payload.phone_number), address=payload.address, slot_id=payload.slot_id, guest_count=payload.guest_count)
    db.add(registration)
    try:
        db.flush()
        db.add(ExportOutbox(registration_id=registration.id))
        db.commit()
    except IntegrityError:
        # Two requests can pass the initial lookup at the same time. The unique
        # idempotency constraint makes one winner; return that committed booking.
        db.rollback()
        existing = db.scalar(select(Registration).where(Registration.idempotency_key == payload.idempotency_key).options(joinedload(Registration.slot)))
        if existing:
            return existing
        raise
    db.refresh(registration)
    registration.slot = slot
    return registration


def safe_cell(value: str) -> str:
    value = str(value)
    return "'" + value if value[:1] in {"=", "+", "-", "@"} else value


def excel_datetime(value: datetime) -> datetime:
    """Excel stores wall-clock datetimes without timezone metadata."""
    return value.replace(tzinfo=None) if value.tzinfo else value


def write_export(db: Session, path: str) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)
    registrations = db.scalars(select(Registration).where(Registration.status == RegistrationStatus.CONFIRMED).options(joinedload(Registration.slot)).order_by(Registration.created_at)).all()
    slots = db.scalars(select(TimeSlot).order_by(TimeSlot.starts_at)).all()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Registrations"
    headers = ["Confirmation Code", "Registration Timestamp", "Full Name", "Phone Number", "Address", "Slot Label", "Slot Start", "Slot End", "Guest Count", "Registration Status"]
    sheet.append(headers)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for item in registrations:
        sheet.append([item.confirmation_code, excel_datetime(item.created_at), safe_cell(item.full_name), safe_cell(item.phone_number), safe_cell(item.address), item.slot.label, excel_datetime(item.slot.starts_at), excel_datetime(item.slot.ends_at), item.guest_count, item.status])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column, width in {"A": 18, "B": 23, "C": 24, "D": 18, "E": 38, "F": 24, "G": 22, "H": 22, "I": 12, "J": 20}.items():
        sheet.column_dimensions[column].width = width
    summary = workbook.create_sheet("Slot Summary")
    summary.append(["Slot", "Start", "End", "Capacity", "Booked Guests", "Remaining Guests", "Registrations", "Status"])
    for cell in summary[1]:
        cell.font = Font(bold=True)
    counts = dict(db.execute(select(Registration.slot_id, func.count(Registration.id)).where(Registration.status == RegistrationStatus.CONFIRMED).group_by(Registration.slot_id)).all())
    for slot in slots:
        summary.append([slot.label, excel_datetime(slot.starts_at), excel_datetime(slot.ends_at), slot.capacity, slot.booked_guests, max(0, slot.capacity - slot.booked_guests), counts.get(slot.id, 0), slot.status])
    summary.freeze_panes = "A2"
    summary.auto_filter.ref = summary.dimensions
    info = workbook.create_sheet("Export Info")
    info.append(["Field", "Value"])
    info.append(["Generated At", excel_datetime(datetime.now(timezone.utc))])
    info.append(["Event", get_settings().event_name])
    info.append(["Source", "PostgreSQL snapshot"])
    fd, temp_name = tempfile.mkstemp(prefix="registrations-", suffix=".xlsx", dir=output.parent)
    os.close(fd)
    try:
        workbook.save(temp_name)
        load_workbook(temp_name, read_only=True).close()
        os.replace(temp_name, output)
    finally:
        if os.path.exists(temp_name):
            os.unlink(temp_name)
    return output

