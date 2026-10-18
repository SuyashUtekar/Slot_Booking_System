from datetime import datetime, timedelta, timezone

from openpyxl import load_workbook

from app.db import SessionLocal
from app.services import write_export

def create_slot(client, capacity=10):
    start = datetime.now(timezone.utc) + timedelta(days=1)
    response = client.post("/api/v1/admin/slots", headers={"Authorization": "Bearer test-token"}, json={"label": "Morning", "starts_at": start.isoformat(), "ends_at": (start + timedelta(hours=1)).isoformat(), "capacity": capacity})
    assert response.status_code == 201
    return response.json()["id"]

def payload(slot_id, guests=2, key="idempotency-key-001"):
    return {"full_name": "Asha Rao", "phone_number": "+91 98765 43210", "address": "12 Main Street", "slot_id": slot_id, "guest_count": guests, "idempotency_key": key}

def test_booking_consumes_guest_capacity(client):
    slot = create_slot(client, 5)
    response = client.post("/api/v1/registrations", json=payload(slot, 3))
    assert response.status_code == 201
    assert client.get("/api/v1/slots").json()["slots"][0]["remaining_guests"] == 2

def test_booking_rejects_over_capacity_without_consuming(client):
    slot = create_slot(client, 3)
    response = client.post("/api/v1/registrations", json=payload(slot, 4))
    assert response.status_code == 409
    assert client.get("/api/v1/slots").json()["slots"][0]["booked_guests"] == 0

def test_idempotency_returns_original_booking(client):
    slot = create_slot(client, 10)
    first = client.post("/api/v1/registrations", json=payload(slot, 2)).json()
    second = client.post("/api/v1/registrations", json=payload(slot, 2)).json()
    assert second["confirmation_code"] == first["confirmation_code"]
    assert client.get("/api/v1/slots").json()["slots"][0]["booked_guests"] == 2

def test_admin_can_change_guest_limit(client):
    response = client.patch("/api/v1/admin/settings", headers={"Authorization": "Bearer test-token"}, json={"max_guests_per_registration": 10, "count_organizer_as_guest": True})
    assert response.status_code == 200
    slot = create_slot(client, 20)
    too_many = client.post("/api/v1/registrations", json=payload(slot, 11, "idempotency-key-002"))
    assert too_many.status_code == 422


def test_excel_export_accepts_database_datetimes(client, tmp_path):
    slot = create_slot(client, 10)
    assert client.post("/api/v1/registrations", json=payload(slot, 2, "idempotency-key-export")).status_code == 201
    db = SessionLocal()
    try:
        output = write_export(db, str(tmp_path / "registrations.xlsx"))
    finally:
        db.close()
    workbook = load_workbook(output, read_only=True)
    assert {"Registrations", "Slot Summary", "Export Info"}.issubset(workbook.sheetnames)
    workbook.close()

